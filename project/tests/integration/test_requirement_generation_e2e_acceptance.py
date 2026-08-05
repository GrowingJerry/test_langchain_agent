from __future__ import annotations

from pathlib import Path

import pytest

import application.services.project_service as manager_module
from application.services.generation_service import GenerationRequest
from application.services.test_type_recommendation import (
    apply_manual_case_type,
    sync_case_type_state,
)
from application.services.ui_application_service import UIApplicationService
from config.settings import Settings
from infrastructure.exporters.project_documents import build_project_export_rows
from application.services.project_service import ProjectManager


def test_formal_srs_requirement_to_generation_acceptance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(manager_module, "PROJECT_OUTPUT_ROOT", tmp_path / "projects")
    monkeypatch.setattr(
        "application.services.generation_context.search_project_chunks",
        lambda manager, project_id, query, top_k: manager.list_chunks(project_id, 100),
    )
    manager = ProjectManager(tmp_path / "workspace.db")
    service = UIApplicationService(
        manager,
        settings=Settings(enable_agent=False, enable_ollama=False),
    )
    project_id = manager.create_project("验收项目")["project_id"]
    sample = Path(__file__).resolve().parents[2] / "data" / "XX系统软件需求规格说明.docx"
    target = manager.uploads_dir(project_id) / sample.name
    target.write_bytes(sample.read_bytes())
    document_id = manager.add_document(project_id, target.name, "docx", target)
    chunk_id = manager.replace_chunks(
        project_id,
        document_id,
        [{"chunk_text": "REQ-FUNC-001 态势显示功能 系统应实时显示态势信息。", "source_type": "docx"}],
    )[0]
    manager.save_profile(project_id, {"project_name": "验收项目", "test_object": "XX系统"})

    preview = service.preview_requirement_extraction(project_id)
    report = preview["report"]
    machine_rows = preview["requirements"]
    assert report["requirement_count"] >= 10
    assert report["original_id_count"] >= 1
    assert report["filtered_template_instruction_count"] >= 1
    assert report["recommended_test_type_count"] == report["requirement_count"]

    req = next(row for row in machine_rows if row["requirement_id"] == "REQ-FUNC-001")
    assert "4.3.1" in " / ".join(req["section_path"])
    assert req["recommended_test_type"]
    assert req["source_evidence"][0]["block_id"]
    assert "REQ-FUNC-001" in req["description"]

    req["machine_extraction"] = dict(req)
    reviewed = service.save_reviewed_requirements(project_id, [req], machine_rows)
    assert reviewed[0]["machine_extraction"]["requirement_id"] == "REQ-FUNC-001"
    saved = manager.get_requirement(project_id, "REQ-FUNC-001")
    assert saved and saved["source_evidence"]

    state: dict[str, object] = {}
    recommendation = sync_case_type_state(state, [saved], current_default="功能测试")
    assert recommendation["selected_case_type"] == saved["recommended_test_type"]
    apply_manual_case_type(state, "性能测试")
    recommendation = state["case_type_recommendation"]
    assert recommendation["selected_case_type"] == "性能测试"
    assert recommendation["case_type_manually_overridden"] is True

    manager.replace_full_scenario_cards(
        project_id,
        [{
            "scenario_id": "SCN-REQ-FUNC-001",
            "scenario_name": "态势显示验收",
            "related_requirements": ["REQ-FUNC-001"],
            "preconditions": ["系统已启动"],
            "input_data": ["态势数据"],
            "normal_flow": ["显示态势"],
            "source_chunk_ids": [chunk_id],
        }],
    )
    result = service.generate(
        GenerationRequest(
            project_id=project_id,
            requirement_ids=["REQ-FUNC-001"],
            scenario_ids=["SCN-REQ-FUNC-001"],
            case_type=recommendation["selected_case_type"],
            case_count=2,
            requested_mode="rule",
            recommended_test_type=recommendation["recommended_case_type"],
            selected_test_type=recommendation["selected_case_type"],
            test_type_overridden=True,
            test_type_confidence=float(recommendation["recommendation_confidence"]),
            test_type_reasons=list(recommendation["recommendation_reasons"]),
        )
    )
    assert len(result.cases) == 2
    assert all(case.persistence_data["case_type"] == "性能测试" for case in result.cases)
    assert all("REQ-FUNC-001" in case.case.requirement_ids for case in result.cases)
    assert manager.list_trace_sources(project_id)
    matrix = build_project_export_rows(manager, project_id)
    assert matrix["requirements"]
    assert matrix["test_cases"]


def test_template_documents_do_not_become_large_fake_requirement_sets(tmp_path: Path) -> None:
    from docx import Document
    from workflows.learning.structured_requirement_extractor import (
        extract_requirements_from_blocks,
        parse_docx_blocks,
    )

    path = tmp_path / "用户需求说明书模板.docx"
    doc = Document()
    doc.add_heading("用户需求说明书", level=1)
    doc.add_paragraph("本章应描述用户需求。")
    doc.add_paragraph("在形成最后文档时，应删除文档中的所有注释。")
    doc.add_paragraph("XX系统应在此填写功能，例如XXXX。")
    doc.save(path)
    rows = extract_requirements_from_blocks(parse_docx_blocks(path, "TPL"))
    assert rows == []
