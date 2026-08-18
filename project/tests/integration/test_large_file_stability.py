from __future__ import annotations
from pathlib import Path
from docx import Document
from application.services.project_service import ProjectManager
from application.services.ui_application_service import UIApplicationService
from config.settings import Settings

def _formal_docx(path:Path)->Path:
    doc=Document(); doc.add_heading("3.2 子系统总体需求",1); doc.add_heading("3.2.1 [PORTAL] 门户",2); doc.add_paragraph("门户总体说明")
    doc.add_heading("3.3 详细功能需求",1); doc.add_heading("3.3.1 [PORTAL] 门户",2); doc.add_heading("3.3.1.1 功能需求",3); doc.add_heading("3.3.1.1.1 [PROFILE] 个人信息配置",4)
    for title,text in (("功能描述","用户可以查看并修改个人信息。"),("输入","姓名、手机号和邮箱。"),("处理","校验输入格式并保存。"),("输出","显示保存结果。")):
        doc.add_heading(title,5); doc.add_paragraph(text)
    for i in range(300): doc.add_paragraph(f"附加说明 {i}：此段用于模拟接近正式文档的多层正文。")
    doc.save(path); return path

def _large_html()->bytes:
    row='<label for="field-{0}">字段{0}</label><input id="field-{0}" name="field-{0}" value="{1}">'
    body="".join(row.format(i,"x"*100) for i in range(18000))
    data=("<!doctype html><html><head><title>大型配置页</title><link rel='stylesheet' href='app.css'></head><body><form>"+body+"<button>保存</button></form><script src='app.js'></script></body></html>").encode()
    assert len(data)>=2_900_000
    return data

def test_large_docx_and_html_are_bounded_and_playwright_independent(tmp_path,monkeypatch):
    manager=ProjectManager(tmp_path/"workspace.db"); project_id=manager.create_project("稳定性")['project_id']
    service=UIApplicationService(manager,None,Settings(enable_ollama=False))
    progress=[]; parsed=service.analyze_csci_docx(project_id,_formal_docx(tmp_path/"formal.docx"))
    assert parsed["testable_count"]>0 and parsed["section_32_count"]>0 and parsed["section_33_count"]>0
    monkeypatch.setitem(__import__('sys').modules,'playwright',None)
    summary=service.analyze_offline_html(project_id,"large.html",_large_html(),include_elements=False)
    assert "elements" not in summary and summary["element_count"]>10000
    page=service.list_html_elements_page(project_id,1,100)
    assert page["total"]==summary["element_count"] and len(page["rows"])==100
    assert service.workflow_status(project_id)["lowest_function_count"]>0

def test_empty_tree_blocks_atomization_state(tmp_path):
    manager=ProjectManager(tmp_path/"workspace.db"); project_id=manager.create_project("空树")['project_id']
    service=UIApplicationService(manager,None,Settings(enable_ollama=False)); doc=Document(); doc.add_paragraph("没有正式章节"); path=tmp_path/"empty.docx"; doc.save(path)
    parsed=service.analyze_csci_docx(project_id,path)
    assert parsed["node_count"]==0 and parsed["warnings"]
    assert service.workflow_status(project_id)["lowest_function_count"]==0
