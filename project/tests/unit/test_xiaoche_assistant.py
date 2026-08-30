from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict

from application.assistant.intent_router import page_url, route_intent
from application.assistant.tool_registry import ToolRegistry
from desktop_pet.notifications import NotificationLedger
from desktop_pet.status import map_status
from infrastructure.assistant_exporters import AssistantExporter
from infrastructure.documents.assistant_ingestor import AssistantIngestor, safe_filename


def test_document_type_and_path_safety(tmp_path: Path) -> None:
    assert safe_filename("../需求.txt") == "需求.txt"
    ingestor = AssistantIngestor(tmp_path / "work", max_bytes=5)
    too_large = tmp_path / "large.txt"; too_large.write_bytes(b"123456")
    with pytest.raises(ValueError, match="大小限制"): ingestor.ingest(too_large, "s1")
    empty = tmp_path / "empty.md"; empty.touch()
    with pytest.raises(ValueError, match="空文件"): ingestor.ingest(empty, "s1")


def test_docx_and_xlsx_ingestion(tmp_path: Path) -> None:
    docx = tmp_path / "input.docx"; doc = Document(); doc.add_heading("需求标题", 1); doc.add_paragraph("需求正文"); table = doc.add_table(rows=1, cols=2); table.cell(0,0).text="编号"; table.cell(0,1).text="REQ-1"; doc.save(docx)
    parsed = AssistantIngestor(tmp_path / "work").ingest(docx, "session")
    assert "需求标题" in parsed.text and parsed.metadata["tables"][0][0][1] == "REQ-1"
    xlsx = tmp_path / "input.xlsx"; wb=Workbook(); wb.active.title="需求"; wb.active.append(["编号","公式"]); wb.active.append(["R1","=1+1"]); wb.create_sheet("状态").append(["完成"]); wb.save(xlsx)
    parsed_xlsx = AssistantIngestor(tmp_path / "work").ingest(xlsx, "session")
    assert parsed_xlsx.metadata["sheet_names"] == ["需求", "状态"]
    assert "=1+1" in parsed_xlsx.text


def test_docx_xlsx_and_text_export_readback(tmp_path: Path) -> None:
    exporter = AssistantExporter(tmp_path / "outputs")
    docx = exporter.create_docx("s1", "报告", {"title":"测试报告", "paragraphs":["实际内容"], "tables":[[["编号","结果"],["T1","通过"]]]})
    reopened = Document(docx.absolute_path); assert "实际内容" in "".join(p.text for p in reopened.paragraphs); assert reopened.tables[0].cell(1,0).text == "T1"
    xlsx = exporter.create_xlsx("s1", "用例", [{"name":"用例", "rows":[["编号","结果"],["T1","通过"]]}])
    workbook = load_workbook(xlsx.absolute_path); assert workbook["用例"]["B2"].value == "通过"; workbook.close()
    txt = exporter.create_txt("s1", "说明", "中文内容"); assert Path(txt.absolute_path).read_text(encoding="utf-8") == "中文内容"
    with pytest.raises(ValueError): exporter.create_markdown("s1", "空", "")


def test_pdf_export_readback_when_available(tmp_path: Path) -> None:
    pytest.importorskip("fitz")
    exported = AssistantExporter(tmp_path / "outputs").create_pdf("s1", "报告", "中文测试报告\n正文不为空")
    import fitz
    document = fitz.open(exported.absolute_path); assert document.page_count >= 1; assert "报告" in "".join(page.get_text() for page in document); document.close()


class Args(BaseModel):
    model_config = ConfigDict(extra="forbid")
    value: int


def test_registry_whitelist_validation_and_audit(tmp_path: Path) -> None:
    registry = ToolRegistry(tmp_path / "audit.jsonl"); registry.register("double", Args, lambda value: value * 2)
    assert registry.execute("s", "double", {"value": 2}).success
    assert not registry.execute("s", "shell", {}).success
    assert not registry.execute("s", "double", {"value": "bad", "extra": 1}).success
    assert len((tmp_path / "audit.jsonl").read_text(encoding="utf-8").splitlines()) == 3


def test_status_notification_and_navigation(tmp_path: Path) -> None:
    assert map_status("parsing") == "reading"; assert map_status("review_pending") == "reviewing"; assert map_status("unknown") == "idle"
    ledger = NotificationLedger(tmp_path / "seen.json"); assert ledger.should_notify("run-1"); ledger.mark_read("run-1"); assert not ledger.should_notify("run-1"); assert not NotificationLedger(tmp_path / "seen.json").should_notify("run-1")
    intent = route_intent("打开用例审查页面"); assert intent.tool == "open_streamlit_page" and intent.arguments["page"] == "用例审查"
    assert "page=" in page_url("http://127.0.0.1:8501", "导出中心")
    with pytest.raises(ValueError): page_url("http://127.0.0.1:8501", "第七页")
