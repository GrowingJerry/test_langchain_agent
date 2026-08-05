from __future__ import annotations

import sqlite3
from pathlib import Path

from docx import Document

import application.services.project_service as manager_module
from application.services.project_service import ProjectManager
from infrastructure.database.connection import SQLiteConnectionManager
from infrastructure.database.migrations import migrate_database
from workflows.learning.structured_requirement_extractor import (
    extract_requirements_from_blocks,
    extract_structured_requirements,
    infer_test_type,
    parse_docx_blocks,
)


def _docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("4 软件需求", level=1)
    doc.add_heading("4.3 功能需求", level=2)
    doc.add_paragraph("REQ-FUNC-001 态势显示功能")
    doc.add_paragraph("输入：目标位置、航速和航向")
    doc.add_paragraph("处理规则：系统应融合多源目标数据并刷新态势图")
    doc.add_paragraph("输出：显示目标列表和态势图层")
    doc.add_paragraph("异常处理：数据无效时应提示并保持上一帧")
    doc.add_paragraph("测试类型：功能测试")
    doc.add_heading("4.4 性能需求", level=2)
    doc.add_paragraph("REQ-PERF-001 系统端到端延迟应≤100毫秒，连续运行24小时。")
    doc.add_heading("4.5 安全需求", level=2)
    doc.add_paragraph("REQ-SEC-001 系统必须支持登录权限控制，并对传输数据加密。")
    doc.add_heading("4.6 接口需求", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "需求编号"
    table.rows[0].cells[1].text = "需求名称"
    table.rows[0].cells[2].text = "接口参数"
    table.rows[0].cells[3].text = "测试类型"
    row = table.add_row().cells
    row[0].text = "REQ-IF-001"
    row[1].text = "目标上报接口"
    row[2].text = "报文格式应包含id、time、position字段"
    row[3].text = "1"
    doc.add_paragraph("注：本条应说明如何填写接口需求。")
    doc.save(path)


def test_docx_blocks_keep_heading_and_table_relationships(tmp_path: Path) -> None:
    path = tmp_path / "srs.docx"
    _docx(path)
    blocks = parse_docx_blocks(path, "DOC-1")
    assert any(b.block_type == "heading" and b.heading_level == 2 and "功能需求" in b.text for b in blocks)
    table_row = next(b for b in blocks if b.block_type == "table_row")
    assert table_row.table_index == 1
    assert table_row.row_index == 2
    assert "需求编号: REQ-IF-001" in table_row.text
    assert "接口参数: 报文格式" in table_row.text


def test_requirements_preserve_original_ids_and_group_structured_fields(tmp_path: Path) -> None:
    path = tmp_path / "srs.docx"
    _docx(path)
    rows = extract_requirements_from_blocks(parse_docx_blocks(path, "DOC-1"))
    by_id = {row["requirement_id"]: row for row in rows}
    assert "REQ-FUNC-001" in by_id
    assert by_id["REQ-FUNC-001"]["inputs"] == ["目标位置、航速和航向"]
    assert by_id["REQ-FUNC-001"]["processing_rules"]
    assert by_id["REQ-FUNC-001"]["outputs"]
    assert by_id["REQ-FUNC-001"]["exception_rules"]
    assert by_id["REQ-FUNC-001"]["recommended_test_type"] == "功能测试"
    assert by_id["REQ-PERF-001"]["recommended_test_type"] == "性能测试"
    assert by_id["REQ-SEC-001"]["recommended_test_type"] == "安全性测试"
    assert by_id["REQ-IF-001"]["recommended_test_type"] == "接口测试"
    assert by_id["REQ-IF-001"]["need_human_confirm"] is True
    assert "本条应" not in "\n".join(row["description"] for row in rows)


def test_blank_template_does_not_create_fake_requirements(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("用户需求说明书", level=1)
    doc.add_paragraph("本章应描述用户需求。")
    doc.add_paragraph("XX系统应在此填写功能，例如XXXX。")
    path = tmp_path / "template.docx"
    doc.save(path)
    assert extract_requirements_from_blocks(parse_docx_blocks(path, "DOC-T")) == []


def test_test_type_inference_multiple_candidates_and_numeric_confirm() -> None:
    verdict = infer_test_type(
        {
            "requirement_id": "REQ-PERF-009",
            "title": "接口吞吐",
            "description": "接口报文吞吐量不低于1000条/秒，故障后应恢复。",
            "section_path": ["4 性能需求"],
            "performance_constraints": ["吞吐量不低于1000条/秒"],
        }
    )
    assert verdict["recommended_test_type"] == "性能测试"
    assert "接口测试" in verdict["alternative_test_types"]
    assert "可靠性测试" in verdict["alternative_test_types"]
    numeric = infer_test_type({"requirement_id": "REQ-1", "title": "登录", "description": "登录权限"}, "2")
    assert numeric["need_human_confirm"] is True


def test_old_database_requirement_rows_remain_readable(tmp_path: Path) -> None:
    db_path = tmp_path / "old.db"
    conn = sqlite3.connect(db_path)
    conn.executescript("""
        CREATE TABLE projects(project_id TEXT PRIMARY KEY, project_name TEXT NOT NULL, description TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE project_requirements(requirement_id TEXT NOT NULL, project_id TEXT NOT NULL, title TEXT, description TEXT, category TEXT, source_document TEXT, source_chunk_id TEXT, created_at TEXT NOT NULL, PRIMARY KEY(project_id, requirement_id));
        INSERT INTO projects VALUES('P1','旧项目','','2020','2020');
        INSERT INTO project_requirements VALUES('REQ-OLD','P1','旧需求','系统应支持查询','功能需求','old.txt','CHK-1','2020');
    """)
    conn.commit()
    conn.close()
    migrate_database(SQLiteConnectionManager(db_path))
    manager = ProjectManager(db_path)
    row = manager.get_requirement("P1", "REQ-OLD")
    assert row
    assert row["source_chunk_ids"] == ["CHK-1"]
    assert row["recommended_test_type"] == ""


def test_txt_documents_use_legacy_fallback(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(manager_module, "PROJECT_OUTPUT_ROOT", tmp_path / "projects")
    manager = ProjectManager(tmp_path / "workspace.db")
    project_id = manager.create_project("TXT")["project_id"]
    path = tmp_path / "req.txt"
    path.write_text("系统应支持用户登录。\n接口应支持JSON报文。", encoding="utf-8")
    doc_id = manager.add_document(project_id, path.name, "txt", path)
    manager.replace_chunks(project_id, doc_id, [{"text": path.read_text(encoding="utf-8")}])
    rows = extract_structured_requirements(manager, project_id)
    assert rows
    assert rows[0]["requirement_id"].startswith("REQ-")
