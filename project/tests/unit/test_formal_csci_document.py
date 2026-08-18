from io import BytesIO
import base64
from docx import Document
from docx.shared import Inches
from application.services.csci_document_service import parse_formal_csci_docx

PNG=base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z4m0AAAAASUVORK5CYII=")

def _build(path):
    doc=Document(); doc.add_heading("3.2 CSCI能力需求",1); doc.add_heading("3.2.1 ZH_TYMH统一门户",2); doc.add_paragraph("统一提供门户能力。")
    doc.add_heading("3.3 CSCI能力",1); doc.add_heading("3.3.3 ZH_TYMH统一门户",2); doc.add_heading("3.3.3.1 功能需求",3)
    doc.add_heading("3.3.3.1.1 ZH_TYMH_XWMH新闻门户",4); doc.add_heading("3.3.3.1.1.1 功能描述",5); doc.add_paragraph("支持新增和编辑公告。")
    doc.add_heading("3.3.3.1.1.2 输入",5); table=doc.add_table(rows=2,cols=2); table.cell(0,0).text="字段"; table.cell(0,1).text="规则"; table.cell(1,0).text="标题"; table.cell(1,1).text="必填"
    doc.add_paragraph("输入结束。")
    doc.add_heading("3.3.3.1.1.3 处理",5); p=doc.add_paragraph("流程图1 公告处理流程"); p.add_run().add_picture(BytesIO(PNG),width=Inches(.1))
    doc.add_heading("3.3.3.1.1.4 输出",5); doc.add_paragraph("显示本地保存提示。")
    doc.add_heading("3.4 外部接口",1); doc.add_paragraph("不得混入。")
    doc.save(path)

def test_formal_32_33_tree_and_ordered_evidence(tmp_path):
    path=tmp_path/"formal.docx"; _build(path); result=parse_formal_csci_docx(path)
    overview=result["overview_nodes"][0]; leaf=result["testable_nodes"][0]
    assert overview.identifier=="ZH_TYMH" and overview.name=="统一门户"
    assert leaf.identifier=="ZH_TYMH_XWMH" and leaf.overview_node_id==overview.node_id
    assert leaf.hierarchy_path==["统一门户","功能需求","新闻门户"]
    assert "标题 | 必填" in leaf.sections["输入"] and "输入结束" in leaf.sections["输入"]
    kinds=[x["kind"] for x in leaf.section_evidence["处理"]]
    assert "paragraph" in kinds and "image" in kinds
    assert "不得混入" not in str(result)
