from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from application.assistant.gateway import AssistantGateway
from application.services.project_service import ProjectManager
from config.settings import Settings


class FakeUIService:
    def __init__(self,manager): self.manager=manager
    def list_projects(self): return self.manager.list_projects()


def requirement_docx(path:Path)->Path:
    doc=Document(); doc.core_properties.title="接口需求说明书"; doc.add_heading("接口需求",0)
    doc.add_heading("功能需求",1); doc.add_paragraph("REQ-001：系统应支持用户登录。")
    doc.add_paragraph("REQ-002：系统必须记录登录审计日志。")
    doc.add_paragraph("系统应在网络中断时给出明确提示。")
    table=doc.add_table(rows=2,cols=2); table.cell(0,0).text="需求标识"; table.cell(0,1).text="描述"; table.cell(1,0).text="REQ-003"; table.cell(1,1).text="系统应支持退出登录。"; doc.save(path); return path


@pytest.fixture()
def gateway(tmp_path:Path)->AssistantGateway:
    manager=ProjectManager(tmp_path/"project.db")
    settings=Settings(xiaoche_output_dir="outputs/assistant")
    return AssistantGateway(FakeUIService(manager),settings,tmp_path,chat_handler=lambda text,history:f"OLLAMA:{text}")


def test_attachment_id_enters_message_and_does_not_leak_to_next(gateway:AssistantGateway,tmp_path:Path)->None:
    record=gateway.upload_attachment("s1",requirement_docx(tmp_path/"requirements.docx")); assert record["attachment_id"] and record["parse_status"]=="pending"
    response=gateway.handle_message("s1","","这是什么",[record["attachment_id"]]); assert "Word" not in response.message or "DOCX" in response.message; assert "接口需求说明书" in response.message; assert "小测" not in response.message
    messages=gateway.sessions.page("s1"); assert record["attachment_id"] in messages[0]["attachment_ids_json"]
    ordinary=gateway.handle_message("s1","","什么是边界值测试？",[]); assert ordinary.message.startswith("OLLAMA:")
    assert gateway.sessions.page("s1")[-2]["attachment_ids_json"]=="[]"


def test_parse_document_real_metadata_and_persistence(gateway:AssistantGateway,tmp_path:Path)->None:
    record=gateway.upload_attachment("s2",requirement_docx(tmp_path/"requirements.docx")); response=gateway.handle_message("s2","","帮我解析这个文件",[record["attachment_id"]])
    assert response.status=="success" and response.tool_runs[0]["tool"]=="parse_document"
    parsed=gateway.sessions.get_attachment("s2",record["attachment_id"]); assert parsed["parse_status"]=="success"; assert parsed["paragraph_count"]>=4; assert parsed["table_count"]==1; assert parsed["parsed_json"]["metadata"]["tables"][0][1][0]=="REQ-003"
    reopened=AssistantGateway(FakeUIService(gateway.service.manager),gateway.settings,gateway.project_root,chat_handler=lambda text,history:"ok")
    assert reopened.sessions.get_attachment("s2",record["attachment_id"])["parse_status"]=="success"


def test_word_traceability_matrix_end_to_end(gateway:AssistantGateway,tmp_path:Path)->None:
    record=gateway.upload_attachment("s3",requirement_docx(tmp_path/"requirements.docx")); response=gateway.handle_message("s3","","帮我解析这个文件，转化成Word版本的需求追踪矩阵",[record["attachment_id"]])
    assert response.status=="success"; assert [run["tool"] for run in response.tool_runs]==["parse_document","create_docx"]
    assert len(response.output_files)==1; output=Path(response.output_files[0]["absolute_path"]); assert output.is_file(); assert gateway.output_root in output.parents
    doc=Document(output); assert len(doc.tables)==1; headers=[cell.text for cell in doc.tables[0].rows[0].cells]; assert headers==["序号","需求标识","需求名称","需求描述","来源章节/位置","关联测试类型","关联测试用例","覆盖状态","备注"]; assert len(doc.tables[0].rows)>=4
    rows=[[cell.text for cell in row.cells] for row in doc.tables[0].rows[1:]]; assert any(row[1]=="REQ-001" for row in rows); assert any(row[1].startswith("TEMP-") and "待人工确认" in row[8] for row in rows)
    assert gateway.sessions.list_outputs("s3")[0]["validation_status"]=="passed"


def test_corrupt_docx_reports_real_failure(gateway:AssistantGateway,tmp_path:Path)->None:
    bad=tmp_path/"bad.docx"; bad.write_bytes(b"not-a-docx"); record=gateway.upload_attachment("s4",bad); response=gateway.handle_message("s4","","解析这个文件",[record["attachment_id"]])
    assert response.status=="error" and "失败阶段" in response.message
    assert gateway.sessions.get_attachment("s4",record["attachment_id"])["parse_status"]=="failed"


def test_output_path_cannot_escape(gateway:AssistantGateway,tmp_path:Path)->None:
    with pytest.raises(ValueError): gateway.open_output_path(str(tmp_path.parent/"outside.txt"),False)


def test_above_content_word_export_uses_session_context(gateway:AssistantGateway)->None:
    first=gateway.handle_message("s5","","解释边界值测试",[]); assert first.message.startswith("OLLAMA:")
    exported=gateway.handle_message("s5","","把上面的内容生成Word",[]); assert exported.status=="success"; assert exported.tool_runs[0]["tool"]=="create_docx"
    document=Document(exported.output_files[0]["absolute_path"]); assert "OLLAMA:解释边界值测试" in "".join(p.text for p in document.paragraphs)


def test_export_failure_never_claims_success(gateway:AssistantGateway,monkeypatch)->None:
    gateway.handle_message("s6","","先给出一段有效内容",[])
    def fail(**kwargs): raise RuntimeError("disk full")
    monkeypatch.setattr(gateway.exporter,"create_docx",fail)
    # Registry captured the bound method during construction; replace handler explicitly.
    gateway.registry._tools["create_docx"]=(gateway.registry._tools["create_docx"][0],fail)
    response=gateway.handle_message("s6","","生成Word",[]); assert response.status=="error"; assert "已生成" not in response.message
