"""Build the synthetic DOCX fixture. The contents are fictional test data."""
from pathlib import Path
from docx import Document

TARGET = Path(__file__).with_name("synthetic_csci_requirements.docx")


def build(target: Path = TARGET) -> Path:
    doc = Document()
    doc.add_heading("仿真软件需求规格说明书（自动化测试夹具，内容虚构）", 0)
    doc.add_heading("CSCI能力需求", 1)
    for capability, cap_id, page, page_id, function, function_id in (
        ("新闻门户", "XWMP", "个人中心", "GRZX", "个人信息配置", "GRXXPZ"),
        ("系统管理", "XTGL", "用户管理", "YHGL", "用户查询", "YHCX"),
    ):
        doc.add_heading(f"{capability}/标识符_{cap_id}", 2)
        doc.add_heading(f"{page}/标识符_{page_id}", 3)
        doc.add_heading(f"{function}/标识符_{function_id}", 4)
        doc.add_heading("功能描述", 5)
        if function_id == "GRXXPZ":
            doc.add_paragraph("用户可以查看和修改姓名、手机号、电子邮箱。姓名不能为空；手机号应为11位数字；电子邮箱应符合邮箱格式。合法修改后保存信息并提示“保存成功”。")
        else:
            doc.add_paragraph("管理员可以按姓名查询用户，并查看查询结果。")
        doc.add_heading("输入", 5); doc.add_paragraph("页面表单字段和用户操作。")
        doc.add_heading("处理", 5); doc.add_paragraph("校验输入；服务端持久化结果需要联机环境确认。")
        doc.add_heading("输出", 5); doc.add_paragraph("页面显示校验状态或操作提示。")
    doc.add_heading("外部接口", 1); doc.add_paragraph("不属于CSCI能力需求抽取范围。")
    doc.save(target); return target


if __name__ == "__main__":
    print(build())
