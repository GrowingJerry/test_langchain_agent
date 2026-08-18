"""Build fictional formal DOCX and multi-page portal ZIP for real E2E tests."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT=Path(__file__).resolve().parent

def build_docx(path:Path)->Path:
    doc=Document(); doc.add_heading("外网端到端仿真需求（内容虚构，仅供测试）",0)
    doc.add_heading("3.2 CSCI能力需求",1); doc.add_heading("3.2.1 ZH_TYMH统一门户",2); doc.add_paragraph("统一门户提供新闻和公告内容管理入口。")
    doc.add_heading("3.3 CSCI能力",1); doc.add_heading("3.3.3 ZH_TYMH统一门户",2); doc.add_heading("3.3.3.1 功能需求",3); doc.add_heading("3.3.3.1.1 ZH_TYMH_XWMH新闻门户",4)
    doc.add_heading("3.3.3.1.1.1 功能描述",5); doc.add_paragraph("管理员支持查看公告列表、新增公告和编辑公告。配置标题、类型和内容；保存成功后显示“保存成功”。标题不能为空。")
    doc.add_heading("3.3.3.1.1.2 输入",5); table=doc.add_table(rows=4,cols=3)
    for row,values in zip(table.rows,[("字段","控件","规则"),("标题","文本框","必填"),("类型","下拉框","选择通知或新闻"),("内容","编辑区","普通文本")]):
        for cell,value in zip(row.cells,values): cell.text=value
    doc.add_heading("3.3.3.1.1.3 处理",5); doc.add_paragraph("点击新增打开编辑弹窗；本地校验通过后显示模拟保存结果。真实持久化必须联机验证。")
    doc.add_heading("3.3.3.1.1.4 输出",5); doc.add_paragraph("列表页面显示公告；本地页面显示校验提示或“保存成功”。")
    doc.save(path); return path

def build_site(path:Path)->Path:
    files={
      "index.html":"""<!doctype html><meta charset='utf-8'><title>统一门户</title><nav><a href='news.html'>新闻门户</a></nav>""",
      "news.html":"""<!doctype html><meta charset='utf-8'><title>新闻门户</title><h1>公告列表</h1><button aria-label='新增公告' onclick=\"document.querySelector('#editor').hidden=false\">新增公告</button><div id='editor' role='dialog' aria-label='公告编辑' hidden><label for='title'>标题</label><input id='title' required><label for='type'>类型</label><select id='type'><option>通知</option><option>新闻</option></select><label for='content'>内容</label><textarea id='content'></textarea><button aria-label='保存' onclick=\"const t=document.querySelector('#title');document.querySelector('#msg').textContent=t.value?'保存成功':'标题不能为空'\">保存</button><button aria-label='发布'>发布</button><div id='msg' role='status'></div></div><script src='app.js'></script>""",
      "app.js":"console.log('offline fixture loaded')",
    }
    with ZipFile(path,"w") as zf:
        for name,text in files.items(): zf.writestr(name,text)
    return path

if __name__=="__main__":
    print(build_docx(ROOT/"external_e2e_requirements.docx")); print(build_site(ROOT/"external_e2e_site.zip"))
