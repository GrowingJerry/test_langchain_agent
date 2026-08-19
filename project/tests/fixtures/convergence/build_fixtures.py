"""Build fictional convergence fixtures; no production or personal data."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT=Path(__file__).resolve().parent

def csci():
    d=Document(); d.add_heading('虚构测试夹具：移位CSCI需求',0); d.add_heading('4.2 CSCI能力需求',1); d.add_heading('[DEMO_PORTAL] 示例门户',2); d.add_paragraph('示例门户提供资料维护能力。')
    d.add_heading('4.3 CSCI能力',1); d.add_heading('[DEMO_PORTAL] 示例门户',2); d.add_heading('功能需求',3); d.add_heading('[DEMO_PROFILE] 资料维护',4)
    for title,text in [('功能描述','用户可以查看资料，修改姓名和类别，保存后显示“保存成功”；姓名为空时提示“姓名不能为空”。'),('输入','姓名、类别。'),('处理','校验姓名非空并保存；服务端持久化结果待联机确认。'),('输出','显示保存成功或姓名不能为空。')]: d.add_heading(title,5); d.add_paragraph(text)
    d.save(ROOT/'shifted_csci.docx')

def general():
    d=Document(); d.add_heading('虚构测试夹具：普通产品需求',0); d.add_heading('资料中心',1); d.add_paragraph('用户可以进入资料维护页面。'); d.add_heading('编辑资料',2); d.add_paragraph('功能描述：用户可以修改姓名和类别并保存。'); d.add_paragraph('输入：姓名、类别。'); d.add_paragraph('输出：成功时显示保存成功。'); d.add_paragraph('约束：姓名不能为空，长度最多20个字符。'); d.add_paragraph('异常行为：姓名为空时显示姓名不能为空。'); d.add_paragraph('边界行为：姓名长度为20个字符时允许保存，超过20个字符时阻止提交。'); d.save(ROOT/'general_requirements.docx')

def site():
    bloat='<!--'+('irrelevant comment '*120000)+'--><style>'+('body{color:#333}'*30000)+'</style><script>'+('const unused=1;'*30000)+'</script><img src="data:image/png;base64,'+('A'*500000)+'">'
    index='''<!doctype html><meta charset="utf-8"><title>资料中心</title><a href="edit.html">资料维护</a>'''+bloat
    edit='''<!doctype html><meta charset="utf-8"><title>资料维护</title><form id="profile"><label for="name">姓名</label><input id="name" name="name" required maxlength="20"><label for="category">类别</label><select id="category"><option>个人</option><option>组织</option></select><button type="button" onclick="saveProfile()">保存</button></form><table><tr><th>姓名</th><th>类别</th></tr></table><dialog id="ok">保存成功</dialog><script>function saveProfile(){const n=document.querySelector('#name');if(!n.value){n.setCustomValidity('姓名不能为空');return;}document.querySelector('#ok').setAttribute('open','');}</script>'''
    with ZipFile(ROOT/'offline_site.zip','w') as z: z.writestr('index.html',index); z.writestr('edit.html',edit)

if __name__=='__main__': ROOT.mkdir(parents=True,exist_ok=True); csci(); general(); site()
