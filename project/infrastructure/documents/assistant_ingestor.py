"""Safe, bounded local document staging and parsing."""

from __future__ import annotations

import hashlib
import mimetypes
import re
import shutil
import uuid
from dataclasses import asdict, dataclass
from pathlib import Path

SUPPORTED = {".docx", ".pdf", ".xlsx", ".xls", ".txt", ".md", ".markdown"}


@dataclass(frozen=True)
class ParsedAttachment:
    attachment_id: str; name: str; kind: str; size: int; text: str
    metadata: dict; stored_path: str; sha256: str
    def to_dict(self) -> dict: return asdict(self)


def safe_filename(name: str) -> str:
    value = Path(name).name
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    if not value or value in {".", ".."}: raise ValueError("文件名无效")
    return value[:180]


class AssistantIngestor:
    def __init__(self, workspace: Path, max_bytes: int = 20 * 1024 * 1024, max_chars: int = 200_000, pdf_max_pages: int = 200, excel_max_rows: int = 10_000, excel_max_cols: int = 200) -> None:
        self.workspace=workspace.resolve(); self.max_bytes=max_bytes; self.max_chars=max_chars; self.pdf_max_pages=pdf_max_pages; self.excel_max_rows=excel_max_rows; self.excel_max_cols=excel_max_cols

    def stage(self, source: Path, session_id: str) -> ParsedAttachment:
        source=source.resolve()
        if not source.is_file(): raise ValueError("附件不存在或不是文件")
        size=source.stat().st_size
        if size<=0: raise ValueError("拒绝空文件")
        if size>self.max_bytes: raise ValueError("附件超过大小限制")
        suffix=source.suffix.lower()
        if suffix not in SUPPORTED: raise ValueError("不支持的附件类型")
        attachment_id=uuid.uuid4().hex; target_dir=(self.workspace/safe_filename(session_id)/"attachments").resolve()
        if self.workspace not in target_dir.parents: raise ValueError("会话路径越界")
        target_dir.mkdir(parents=True,exist_ok=True); target=target_dir/f"{attachment_id}-{safe_filename(source.name)}"; shutil.copyfile(source,target)
        return ParsedAttachment(attachment_id,source.name,suffix.lstrip("."),size,"",{"parse_status":"pending","media_type":mimetypes.guess_type(source.name)[0] or "application/octet-stream"},str(target),hashlib.sha256(target.read_bytes()).hexdigest())

    def parse_stored(self, path: Path, attachment_id: str, original_name: str) -> ParsedAttachment:
        path=path.resolve()
        if self.workspace!=path and self.workspace not in path.parents: raise ValueError("附件路径超出受控目录")
        if not path.is_file() or path.stat().st_size<=0: raise ValueError("受控附件不存在或为空")
        suffix=Path(original_name).suffix.lower(); text,metadata=self._parse(path,suffix)
        if not text.strip(): raise ValueError("文档没有可读取内容")
        metadata["truncated"]=len(text)>self.max_chars
        return ParsedAttachment(attachment_id,original_name,suffix.lstrip("."),path.stat().st_size,text[:self.max_chars],metadata,str(path),hashlib.sha256(path.read_bytes()).hexdigest())

    def ingest(self, source: Path, session_id: str) -> ParsedAttachment:
        staged=self.stage(source,session_id)
        return self.parse_stored(Path(staged.stored_path),staged.attachment_id,staged.name)

    def _parse(self,path:Path,suffix:str)->tuple[str,dict]:
        if suffix==".docx": return self._docx(path)
        if suffix==".pdf": return self._pdf(path)
        if suffix==".xlsx": return self._xlsx(path)
        if suffix==".xls": raise ValueError("旧版 XLS 需要安全转换器；不会执行宏或公式")
        return self._text(path)

    def _text(self,path:Path)->tuple[str,dict]:
        raw=path.read_bytes()
        for encoding in ("utf-8-sig","utf-8","gb18030","utf-16"):
            try: return raw.decode(encoding),{"parser_name":"text","encoding":encoding}
            except UnicodeDecodeError: continue
        raise ValueError("无法识别文本编码")

    def _docx(self,path:Path)->tuple[str,dict]:
        from docx import Document
        document=Document(path); lines=[]; styles=[]; blocks=[]; headings=[]
        for index,paragraph in enumerate(document.paragraphs,start=1):
            if not paragraph.text.strip(): continue
            style=paragraph.style.name if paragraph.style else ""; lines.append(paragraph.text); styles.append(style)
            block={"kind":"paragraph","index":index,"text":paragraph.text,"style":style,"source":f"段落{index}"}; blocks.append(block)
            if style.lower().startswith(("heading","标题")): headings.append(block)
        tables=[]
        for table_index,table in enumerate(document.tables,start=1):
            rows=[[cell.text for cell in row.cells] for row in table.rows]; tables.append(rows); lines.extend("\t".join(row) for row in rows)
            blocks.append({"kind":"table","index":table_index,"rows":rows,"row_count":len(rows),"column_count":max((len(row) for row in rows),default=0),"source":f"表格{table_index}"})
        headers=[]; footers=[]
        for section in document.sections:
            headers.extend(p.text for p in section.header.paragraphs if p.text.strip()); footers.extend(p.text for p in section.footer.paragraphs if p.text.strip())
        lines.extend(headers+footers); core=document.core_properties
        return "\n".join(lines),{"parser_name":"python-docx","title":core.title or (headings[0]["text"] if headings else ""),"paragraph_count":len(document.paragraphs),"tables":tables,"table_count":len(tables),"styles":styles,"headings":headings,"blocks":blocks,"headers":headers,"footers":footers,"metadata":{"author":core.author or "","subject":core.subject or ""}}

    def _pdf(self,path:Path)->tuple[str,dict]:
        try: import fitz
        except ImportError as exc: raise RuntimeError("解析 PDF 需要 PyMuPDF") from exc
        document=fitz.open(path)
        if document.page_count>self.pdf_max_pages: document.close(); raise ValueError("PDF 页数超过限制")
        pages=[{"page":index+1,"text":page.get_text("text")} for index,page in enumerate(document)]; document.close()
        return "\n".join(f"[第{x['page']}页]\n{x['text']}" for x in pages),{"parser_name":"PyMuPDF","pages":len(pages),"page_count":len(pages),"page_text":pages}

    def _xlsx(self,path:Path)->tuple[str,dict]:
        from openpyxl import load_workbook
        workbook=load_workbook(path,read_only=False,data_only=False,keep_links=False); output=[]; sheets=[]
        for worksheet in workbook.worksheets:
            max_row=min(worksheet.max_row,self.excel_max_rows); max_col=min(worksheet.max_column,self.excel_max_cols); rows=[]
            for row in worksheet.iter_rows(min_row=1,max_row=max_row,max_col=max_col,values_only=False):
                values=[cell.value for cell in row]; rows.append(values); output.append("\t".join("" if value is None else str(value) for value in values))
                if sum(len(item) for item in output)>=self.max_chars: break
            sheets.append({"name":worksheet.title,"rows":rows,"merged_cells":[str(x) for x in worksheet.merged_cells.ranges],"truncated":worksheet.max_row>max_row or worksheet.max_column>max_col})
        workbook.close(); return "\n".join(output),{"parser_name":"openpyxl","sheets":sheets,"sheet_count":len(sheets),"sheet_names":[x["name"] for x in sheets]}
