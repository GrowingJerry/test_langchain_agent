# -*- coding: utf-8 -*-
"""Project-level Excel and Word document export."""

from __future__ import annotations

from datetime import datetime
import json
import re
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment

from application.services.project_service import ProjectManager


def _safe_name(value: str) -> str:
    """Return a filesystem-safe file stem."""
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value or "project").strip()
    return text[:80] or "project"


def _json_loads_list(value: Any) -> List[str]:
    """Parse a JSON list or comma-separated text into a string list."""
    if isinstance(value, list):
        return [str(x) for x in value if str(x)]
    if not value:
        return []
    if isinstance(value, str):
        try:
            data = json.loads(value)
            if isinstance(data, list):
                return [str(x) for x in data if str(x)]
        except json.JSONDecodeError:
            pass
        return [x.strip() for x in value.split(",") if x.strip()]
    return [str(value)]


def _join(value: Any) -> str:
    """Format lists and dictionaries for export cells."""
    if isinstance(value, list):
        return "\n".join(str(x) for x in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


def _case_field(case_json: Dict[str, Any], *names: str) -> Any:
    """Read a generated case field with fallback aliases."""
    for name in names:
        if name in case_json and case_json[name] not in (None, ""):
            return case_json[name]
    return ""


def _write_sheet(
    workbook: Workbook, title: str, rows: List[Dict[str, Any]], first: bool = False
) -> None:
    """Write a list of dictionaries to an Excel sheet."""
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title[:31]
    headers = list(rows[0].keys()) if rows else ["empty"]
    sheet.append(headers)
    for row in rows:
        sheet.append([_join(row.get(header, "")) for header in headers])
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _export_dir(manager: ProjectManager, project_id: str) -> Path:
    """Return the project document export directory."""
    path = manager.project_dir(project_id) / "exports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def build_project_export_rows(
    manager: ProjectManager, project_id: str
) -> Dict[str, List[Dict[str, Any]]]:
    """Collect profile, requirements, cases, matrix, and trace rows for export."""
    project = manager.get_project(project_id) or {}
    profile = manager.get_profile(project_id) or {}
    requirements = manager.list_requirements(project_id)
    cases = manager.list_generated_cases(project_id)
    traces = manager.list_trace_sources(project_id)

    profile_rows = [
        {"field": "project_id", "value": project_id},
        {
            "field": "project_name",
            "value": profile.get("project_name") or project.get("project_name", ""),
        },
        {"field": "domain", "value": profile.get("domain", "")},
        {"field": "test_object", "value": profile.get("test_object", "")},
        {"field": "main_functions", "value": _join(profile.get("main_functions", []))},
        {"field": "interfaces", "value": _join(profile.get("interfaces", []))},
        {
            "field": "quality_attributes",
            "value": _join(profile.get("quality_attributes", [])),
        },
        {"field": "constraints", "value": _join(profile.get("constraints", []))},
    ]

    requirement_rows = [
        {
            "requirement_id": r.get("requirement_id", ""),
            "title": r.get("title", ""),
            "description": r.get("description", ""),
            "category": r.get("category", ""),
            "source_document": r.get("source_document", ""),
            "source_chunk_id": r.get("source_chunk_id", ""),
        }
        for r in requirements
    ]

    case_rows = []
    matrix_rows = []
    for item in cases:
        case_json = item.get("case_json") or {}
        source_chunk_ids = _json_loads_list(item.get("source_chunk_ids"))
        case_id = item.get("case_id", "")
        requirement_id = item.get("requirement_id", "")
        case_rows.append(
            {
                "case_id": case_id,
                "case_name": _case_field(case_json, "case_name", "用例名称"),
                "case_type": item.get("case_type")
                or _case_field(case_json, "case_type", "测试类别"),
                "requirement_id": requirement_id,
                "test_purpose": _case_field(case_json, "test_purpose", "测试目的"),
                "prerequisites": _case_field(case_json, "prerequisites", "前置条件"),
                "test_steps": _join(_case_field(case_json, "test_steps", "测试步骤")),
                "expected_results": _join(
                    _case_field(
                        case_json, "expected_result", "expected_results", "预期结果"
                    )
                ),
                "pass_criteria": _case_field(case_json, "pass_criteria", "判定准则"),
                "need_human_confirmation": _case_field(
                    case_json,
                    "need_human_confirm",
                    "need_human_confirmation",
                    "是否需要人工确认",
                ),
                "source_document": _join(
                    _case_field(
                        case_json, "source_documents", "source_document", "来源文档"
                    )
                ),
                "source_chunk_ids": "\n".join(source_chunk_ids),
                "generation_run_id": item.get("generation_run_id", ""),
            }
        )
        matrix_rows.append(
            {
                "requirement_id": requirement_id,
                "case_id": case_id,
                "case_name": _case_field(case_json, "case_name", "用例名称"),
                "case_type": item.get("case_type")
                or _case_field(case_json, "case_type", "测试类别"),
                "covered": "Y" if case_id and requirement_id else "N",
                "source_chunk_ids": "\n".join(source_chunk_ids),
            }
        )
    covered_requirement_ids = {row.get("requirement_id") for row in matrix_rows}
    for requirement in requirements:
        requirement_id = requirement.get("requirement_id", "")
        if requirement_id not in covered_requirement_ids:
            matrix_rows.append(
                {
                    "requirement_id": requirement_id,
                    "case_id": "",
                    "case_name": "",
                    "case_type": "",
                    "covered": "N",
                    "source_chunk_ids": requirement.get("source_chunk_id", ""),
                }
            )

    trace_rows = []
    chunk_by_id = {
        c.get("chunk_id"): c for c in manager.list_chunks(project_id, limit=5000)
    }
    for t in traces:
        chunk = chunk_by_id.get(t.get("source_chunk_id")) or {}
        trace_rows.append(
            {
                "trace_id": t.get("trace_id", ""),
                "artifact_type": t.get("artifact_type", ""),
                "artifact_id": t.get("artifact_id", ""),
                "source_document": t.get("source_document", ""),
                "source_chunk_id": t.get("source_chunk_id", ""),
                "chunk_index": chunk.get("chunk_index", ""),
                "snippet": str(chunk.get("content") or "")[:500],
            }
        )

    def scoped_rows(sql: str) -> List[Dict[str, Any]]:
        with manager.connections.connection() as conn:
            return [dict(row) for row in conn.execute(sql, (project_id,)).fetchall()]

    requirement_nodes = scoped_rows("SELECT identifier,name,level,parent_id,hierarchy_path_json,sections_json,source_document,need_human_confirm FROM requirement_nodes WHERE project_id=? ORDER BY level,name")
    indicators = scoped_rows("SELECT indicator_id,capability_id,function_id,indicator_text,indicator_type,source_json,rules_json,verification_scope,need_human_confirm FROM requirement_indicators WHERE project_id=? ORDER BY function_id,indicator_id")
    html_elements = scoped_rows("SELECT page_id,element_id,tag,element_type,element_json FROM html_elements WHERE project_id=? ORDER BY page_id,element_id")
    element_links = scoped_rows("SELECT indicator_id,page_id,confirmed_element_id,confidence,reason,status,need_human_confirm,candidates_json FROM requirement_element_links WHERE project_id=? ORDER BY indicator_id")
    atomic_matrix = scoped_rows("SELECT indicator_id,case_id,case_version,step_numbers_json,coverage_type,coverage_status FROM case_indicator_links WHERE project_id=? ORDER BY indicator_id,case_id")
    versions = scoped_rows("SELECT case_id,version_no,parent_version_no,user_feedback,model_name,changed_fields_json,acceptance_status,operator,created_at FROM case_versions WHERE project_id=? ORDER BY case_id,version_no")
    pending = [row for row in indicators if row.get("need_human_confirm")] + [row for row in element_links if row.get("need_human_confirm")]
    conflicts = [row for row in element_links if row.get("status") == "conflict"]
    return {
        "project_profile": profile_rows,
        "requirements": requirement_rows,
        "test_cases": case_rows,
        "requirement_case_matrix": matrix_rows,
        "trace_sources": trace_rows,
        "requirement_hierarchy": requirement_nodes,
        "atomic_requirements": indicators,
        "html_elements": html_elements,
        "requirement_element_links": element_links,
        "atomic_coverage_matrix": atomic_matrix,
        "case_version_history": versions,
        "pending_confirmation": pending,
        "requirement_html_conflicts": conflicts,
    }


def export_project_excel(manager: ProjectManager, project_id: str) -> Path:
    """Export the current project's testing data to a multi-sheet Excel workbook."""
    project = manager.get_project(project_id) or {}
    profile = manager.get_profile(project_id) or {}
    project_name = (
        profile.get("project_name") or project.get("project_name") or project_id
    )
    date_tag = datetime.now().strftime("%Y%m%d")
    path = (
        _export_dir(manager, project_id)
        / f"{_safe_name(project_name)}_测试文档_{date_tag}.xlsx"
    )
    rows = build_project_export_rows(manager, project_id)
    workbook = Workbook()
    _write_sheet(workbook, "项目画像", rows["project_profile"], first=True)
    _write_sheet(workbook, "需求清单", rows["requirements"])
    _write_sheet(workbook, "测试用例", rows["test_cases"])
    _write_sheet(workbook, "需求-用例追踪矩阵", rows["requirement_case_matrix"])
    _write_sheet(workbook, "来源片段追溯表", rows["trace_sources"])
    for key, title in (
        ("requirement_hierarchy", "需求层级表"),
        ("atomic_requirements", "原子需求表"),
        ("html_elements", "HTML元素表"),
        ("requirement_element_links", "需求元素匹配表"),
        ("atomic_coverage_matrix", "原子需求覆盖矩阵"),
        ("case_version_history", "用例版本历史"),
        ("pending_confirmation", "待人工确认事项"),
        ("requirement_html_conflicts", "需求HTML冲突"),
    ):
        _write_sheet(workbook, title, rows[key])
    workbook.save(path)
    return path


def export_project_word(manager: ProjectManager, project_id: str) -> Path:
    """Export a simple Word test description draft for the current project."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "Word 导出需要安装 python-docx，请执行 pip install python-docx。"
        ) from exc

    project = manager.get_project(project_id) or {}
    profile = manager.get_profile(project_id) or {}
    rows = build_project_export_rows(manager, project_id)
    cases = rows["test_cases"]
    requirements = rows["requirements"]
    project_name = (
        profile.get("project_name") or project.get("project_name") or project_id
    )
    date_tag = datetime.now().strftime("%Y%m%d")
    path = (
        _export_dir(manager, project_id)
        / f"{_safe_name(project_name)}_测试说明初稿_{date_tag}.docx"
    )

    doc = Document()
    doc.add_heading(f"{project_name} 测试说明初稿", level=0)
    doc.add_paragraph(f"生成日期：{datetime.now().strftime('%Y-%m-%d')}")

    doc.add_heading("项目概述", level=1)
    doc.add_paragraph(
        profile.get("domain") or project.get("description") or "需人工补充项目概述。"
    )

    doc.add_heading("测试对象", level=1)
    doc.add_paragraph(profile.get("test_object") or "需人工确认测试对象。")

    doc.add_heading("测试范围", level=1)
    functions = profile.get("main_functions") or []
    doc.add_paragraph(_join(functions) or "按当前项目需求清单和上传文档确定。")

    doc.add_heading("测试依据", level=1)
    doc.add_paragraph(
        "当前项目上传文档、项目需求点、来源片段追溯表、历史测试用例格式参考。"
    )

    doc.add_heading("测试环境", level=1)
    constraints = profile.get("constraints") or []
    doc.add_paragraph(
        _join(constraints) or "需人工确认测试环境、版本、数据和外部接口条件。"
    )

    doc.add_heading("测试方法", level=1)
    methods = sorted(
        {str(c.get("case_type") or "") for c in cases if c.get("case_type")}
    )
    doc.add_paragraph(
        "、".join(methods) if methods else "按生成测试用例中的测试类别和方法执行。"
    )

    doc.add_heading("测试用例摘要", level=1)
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    for idx, title in enumerate(
        ["用例编号", "用例名称", "需求编号", "测试类别", "需人工确认"]
    ):
        hdr[idx].text = title
    for case in cases:
        row = table.add_row().cells
        row[0].text = str(case.get("case_id", ""))
        row[1].text = str(case.get("case_name", ""))
        row[2].text = str(case.get("requirement_id", ""))
        row[3].text = str(case.get("case_type", ""))
        row[4].text = str(case.get("need_human_confirmation", ""))

    doc.add_heading("需求覆盖情况", level=1)
    covered = {c.get("requirement_id") for c in cases if c.get("requirement_id")}
    doc.add_paragraph(
        f"需求总数：{len(requirements)}；已关联测试用例的需求数：{len(covered)}。"
    )

    doc.add_heading("需人工确认的问题", level=1)
    manual_cases = [
        c
        for c in cases
        if str(c.get("need_human_confirmation")).lower() in ("true", "1", "是")
    ]
    if manual_cases:
        for case in manual_cases:
            doc.add_paragraph(
                f"{case.get('case_id', '')}：{case.get('pass_criteria') or '需人工确认'}",
                style="List Bullet",
            )
    else:
        doc.add_paragraph("当前导出的测试用例未标记需人工确认项。")

    doc.save(path)
    return path


def export_project_markdown(manager: ProjectManager, project_id: str) -> Path:
    """Export a lightweight Markdown project test document."""
    project = manager.get_project(project_id) or {}
    profile = manager.get_profile(project_id) or {}
    rows = build_project_export_rows(manager, project_id)
    project_name = (
        profile.get("project_name") or project.get("project_name") or project_id
    )
    date_tag = datetime.now().strftime("%Y%m%d")
    path = (
        _export_dir(manager, project_id)
        / f"{_safe_name(project_name)}_测试文档_{date_tag}.md"
    )

    lines = [
        f"# {project_name} 测试文档",
        "",
        f"- 生成日期：{datetime.now().strftime('%Y-%m-%d')}",
        f"- 项目编号：{project_id}",
        "",
        "## 项目画像",
        "",
    ]
    for row in rows["project_profile"]:
        lines.append(f"- {row.get('field', '')}：{row.get('value', '')}")

    lines.extend(["", "## 需求清单", ""])
    if rows["requirements"]:
        for req in rows["requirements"]:
            lines.append(
                f"- {req.get('requirement_id', '')}：{req.get('title') or req.get('description', '')}"
            )
            if req.get("source_chunk_id"):
                lines.append(f"  - 来源片段：{req.get('source_chunk_id')}")
    else:
        lines.append("- 暂无需求。")

    lines.extend(["", "## 测试用例", ""])
    if rows["test_cases"]:
        for case in rows["test_cases"]:
            lines.append(
                f"### {case.get('case_id', '')} {case.get('case_name', '')}".rstrip()
            )
            lines.append("")
            lines.append(f"- 关联需求：{case.get('requirement_id', '')}")
            lines.append(f"- 测试类型：{case.get('case_type', '')}")
            lines.append(f"- 判定准则：{case.get('pass_criteria', '')}")
            lines.append(f"- 来源片段：{case.get('source_chunk_ids', '')}")
            if case.get("test_steps"):
                lines.extend(["", "测试步骤：", ""])
                for step in str(case.get("test_steps", "")).splitlines():
                    if step.strip():
                        lines.append(f"- {step.strip()}")
            if case.get("expected_results"):
                lines.extend(["", "预期结果：", ""])
                for item in str(case.get("expected_results", "")).splitlines():
                    if item.strip():
                        lines.append(f"- {item.strip()}")
            lines.append("")
    else:
        lines.append("- 暂无测试用例。")

    lines.extend(["", "## 来源追溯", ""])
    if rows["trace_sources"]:
        for trace in rows["trace_sources"]:
            lines.append(
                f"- {trace.get('artifact_id', '')} <- {trace.get('source_document', '')} / "
                f"{trace.get('source_chunk_id', '')}"
            )
    else:
        lines.append("- 暂无来源追溯记录。")

    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return path
