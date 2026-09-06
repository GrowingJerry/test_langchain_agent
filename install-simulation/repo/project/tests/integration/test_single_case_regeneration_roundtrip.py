import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from application.services.case_regeneration_service import CaseRegenerationService
from domain.case_schema import feedback_scope


def test_feedback_scope_recognizes_step_detail_request() -> None:
    assert feedback_scope("将测试步骤写得更明确") == "steps"
    assert feedback_scope("测试步骤不变，只修改预期") == "expected_only"
from application.services.project_service import ProjectManager
from config.settings import Settings
from infrastructure.exporters.project_documents import export_project_excel, export_project_markdown, export_project_word
from infrastructure.repositories.traceability_repository import TraceabilityRepository
from scripts.repair_sparse_case_versions import diagnose, repair


class _Response:
    ok=True; status_code=200
    def __init__(self,payload): self.payload=payload
    def json(self): return {"message":{"content":json.dumps(self.payload,ensure_ascii=False)}}


class _Session:
    def __init__(self,payload): self.payload=payload
    def post(self,*args,**kwargs): return _Response(self.payload)


def _case(project_id,case_id="TC-1"):
    return {"case_id":case_id,"project_id":project_id,"case_name":"保存资料","case_type":"功能测试",
        "requirement_id":"REQ-1","requirement_ids":["REQ-1"],"indicator_ids":["IND-1"],"test_purpose":"验证保存",
        "prerequisites":"用户已登录","test_steps":["输入资料","点击保存"],"expected_result":["输入成功","显示保存成功"],
        "pass_criteria":"结果与预期一致","source_chunk_ids":["CH-1"],"source_documents":["req.docx"],
        "requirement_hierarchy_path":["3.2"],"provenance":{"source":"requirement"},"review_status":"ready","quality_issues":[]}


def test_sparse_legacy_model_response_is_merged_accepted_and_exported(tmp_path):
    manager=ProjectManager(tmp_path/"workspace.db"); project_id=manager.create_project("闭环测试")["project_id"]
    first=_case(project_id); second=_case(project_id,"TC-2"); second["case_name"]="查询资料"
    manager.save_generated_case(project_id,first); manager.save_generated_case(project_id,second)
    settings=Settings(enable_ollama=False,ollama_max_retries=0)
    response={"case_id":"TC-1","steps":first["test_steps"],"expected":["输入成功","显示保存成功，需要待联机验证"]}
    service=CaseRegenerationService(manager,settings,session=_Session(response))
    result=service.regenerate(project_id,"TC-1","只修改预期结果，明确说明保存成功信息需要待联机验证，其他字段保持不变。")
    revised=result["revised"]
    assert "steps" not in revised and "expected" not in revised
    assert revised["test_steps"]==first["test_steps"] and revised["test_purpose"]==first["test_purpose"]
    formal_before=next(x for x in manager.list_generated_cases(project_id) if x["case_id"]=="TC-1")["case_json"]
    assert formal_before==first
    repo=TraceabilityRepository(manager.connections); repo.accept_version(project_id,"TC-1",result["version"]["version_no"])
    formal=next(x for x in manager.list_generated_cases(project_id) if x["case_id"]=="TC-1")["case_json"]
    assert formal==repo.list_case_versions(project_id,"TC-1")[-1]["case_json"]
    assert formal["test_purpose"]==first["test_purpose"] and "待联机验证" in formal["expected_result"][-1]

    xlsx=export_project_excel(manager,project_id); docx=export_project_word(manager,project_id); md=export_project_markdown(manager,project_id)
    sheet=load_workbook(xlsx,data_only=True)["测试用例"]; excel_text="\n".join(str(c.value or "") for row in sheet.iter_rows() for c in row)
    word_text="\n".join(p.text for p in Document(docx).paragraphs)+"\n"+"\n".join(c.text for t in Document(docx).tables for row in t.rows for c in row.cells)
    markdown=Path(md).read_text(encoding="utf-8")
    for text in (excel_text,word_text,markdown):
        assert "待联机验证" in text and "输入资料" in text and "查询资料" in text


def test_accept_missing_formal_case_rolls_back_status(tmp_path):
    manager=ProjectManager(tmp_path/"workspace.db"); project_id=manager.create_project("P")["project_id"]
    repo=TraceabilityRepository(manager.connections)
    version=repo.create_case_version(project_id,"missing",{"case_id":"missing","test_steps":["a"],"expected_result":["b"]})
    with pytest.raises(KeyError): repo.accept_version(project_id,"missing",version["version_no"])
    assert repo.list_case_versions(project_id,"missing")[0]["acceptance_status"]=="proposed"


def test_sparse_repair_is_dry_run_then_backs_up_and_creates_version(tmp_path):
    manager=ProjectManager(tmp_path/"workspace.db"); project_id=manager.create_project("P")["project_id"]
    complete=_case(project_id); manager.save_generated_case(project_id,complete); repo=TraceabilityRepository(manager.connections)
    v1=repo.create_case_version(project_id,"TC-1",complete); repo.accept_version(project_id,"TC-1",v1["version_no"])
    sparse={"case_id":"TC-1","project_id":project_id,"steps":complete["test_steps"],"expected":["新预期1","新预期2"]}
    with manager.connections.transaction() as conn:
        conn.execute("INSERT INTO case_versions(project_id,case_id,version_no,parent_version_no,case_json,acceptance_status) VALUES(?,?,?,?,?,'accepted')",(project_id,"TC-1",2,1,json.dumps(sparse,ensure_ascii=False)))
        conn.execute("UPDATE case_versions SET acceptance_status='rejected' WHERE project_id=? AND case_id=? AND version_no=1",(project_id,"TC-1"))
        conn.execute("UPDATE generated_cases SET case_json=? WHERE project_id=? AND case_id=?",(json.dumps(sparse,ensure_ascii=False),project_id,"TC-1"))
    before=(tmp_path/"workspace.db").read_bytes()
    with manager.connections.connection() as conn: report=diagnose(conn,project_id,"TC-1")[0]
    assert report["recoverable"] and report["latest_complete_version"]==1
    assert (tmp_path/"workspace.db").read_bytes()==before
    applied=repair(tmp_path/"workspace.db",project_id,"TC-1")
    assert Path(applied["backup"]).is_file() and applied["new_version"]==3
    versions=repo.list_case_versions(project_id,"TC-1"); assert len(versions)==3 and versions[-1]["acceptance_status"]=="accepted"
    fixed=manager.list_generated_cases(project_id)[0]["case_json"]
    assert fixed["test_purpose"]==complete["test_purpose"] and fixed["expected_result"]==["新预期1","新预期2"]
