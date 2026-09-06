"""Deterministic local output writers with mandatory read-back validation."""

from __future__ import annotations

import os
import re
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class ExportedFile:
    filename: str
    absolute_path: str
    file_type: str
    size: int
    generated_at: str
    source_attachments: list[str]
    tool: str
    validation: str

    def to_dict(self) -> dict:
        return asdict(self)


class AssistantExporter:
    def __init__(self, output_root: Path, font_name: str = "Microsoft YaHei") -> None:
        self.output_root = output_root.resolve()
        self.font_name = font_name

    def _path(self, session_id: str, filename: str, suffix: str) -> Path:
        session = re.sub(r"[^A-Za-z0-9_-]", "_", session_id)[:80]
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(filename).stem).strip(" .") or "小测文档"
        folder = (self.output_root / session).resolve()
        if self.output_root not in folder.parents:
            raise ValueError("输出路径越界")
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{name}{suffix}"
        if path.exists():
            path = folder / f"{name}-{uuid.uuid4().hex[:8]}{suffix}"
        return path

    def _result(self, path: Path, tool: str, sources: list[str] | None = None) -> ExportedFile:
        if not path.is_file() or path.stat().st_size <= 0:
            raise RuntimeError("输出回读校验失败")
        return ExportedFile(path.name, str(path.resolve()), path.suffix.lstrip("."), path.stat().st_size, datetime.now(timezone.utc).isoformat(), sources or [], tool, "passed")

    @staticmethod
    def _require(content: Any) -> None:
        if content is None or (isinstance(content, str) and not content.strip()) or (isinstance(content, (list, dict)) and not content):
            raise ValueError("拒绝生成空内容文件")

    def create_txt(self, session_id: str, filename: str, content: str, sources: list[str] | None = None) -> ExportedFile:
        self._require(content); path = self._path(session_id, filename, ".txt")
        path.write_text(content, encoding="utf-8")
        if path.read_text(encoding="utf-8") != content: raise RuntimeError("TXT 回读不一致")
        return self._result(path, "create_txt", sources)

    def create_markdown(self, session_id: str, filename: str, content: str, sources: list[str] | None = None) -> ExportedFile:
        self._require(content); path = self._path(session_id, filename, ".md")
        path.write_text(content, encoding="utf-8")
        if not path.read_text(encoding="utf-8").strip(): raise RuntimeError("Markdown 回读为空")
        return self._result(path, "create_markdown", sources)

    def create_docx(self, session_id: str, filename: str, content: dict | str, sources: list[str] | None = None) -> ExportedFile:
        from docx import Document
        from docx.oxml.ns import qn
        self._require(content); data = {"title": filename, "paragraphs": [content]} if isinstance(content, str) else content
        path = self._path(session_id, filename, ".docx"); document = Document()
        normal = document.styles["Normal"]; normal.font.name = self.font_name; normal._element.rPr.rFonts.set(qn("w:eastAsia"), self.font_name)
        if data.get("title"): document.add_heading(str(data["title"]), 0)
        for paragraph in data.get("paragraphs", []): document.add_paragraph(str(paragraph))
        for item in data.get("bullets", []): document.add_paragraph(str(item), style="List Bullet")
        for item in data.get("numbered", []): document.add_paragraph(str(item), style="List Number")
        for table_data in data.get("tables", []):
            if not table_data: continue
            table = document.add_table(rows=0, cols=max(len(row) for row in table_data))
            for row_data in table_data:
                cells = table.add_row().cells
                for index, value in enumerate(row_data): cells[index].text = str(value)
        for _ in range(int(data.get("page_breaks", 0))): document.add_page_break()
        if data.get("header"): document.sections[0].header.paragraphs[0].text = str(data["header"])
        if data.get("footer"): document.sections[0].footer.paragraphs[0].text = str(data["footer"])
        document.save(path)
        reopened = Document(path)
        if not any(p.text.strip() for p in reopened.paragraphs) and not reopened.tables: path.unlink(missing_ok=True); raise RuntimeError("DOCX 回读为空")
        return self._result(path, "create_docx", sources)

    def create_xlsx(self, session_id: str, filename: str, sheets: list[dict], sources: list[str] | None = None) -> ExportedFile:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        self._require(sheets); path = self._path(session_id, filename, ".xlsx"); workbook = Workbook(); workbook.remove(workbook.active)
        has_data = False; thin = Side(style="thin", color="B7C9E2")
        for sheet in sheets:
            worksheet = workbook.create_sheet(str(sheet.get("name") or "数据")[:31]); rows = sheet.get("rows") or []
            for row in rows: worksheet.append(list(row)); has_data = has_data or any(value not in (None, "") for value in row)
            if rows:
                worksheet.freeze_panes = "A2"; worksheet.auto_filter.ref = worksheet.dimensions
                for cell in worksheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="4B83D1")
                for column in worksheet.columns:
                    letter = column[0].column_letter; worksheet.column_dimensions[letter].width = min(50, max(10, max(len(str(c.value or "")) for c in column) + 2))
                for row in worksheet.iter_rows():
                    for cell in row: cell.border = Border(left=thin, right=thin, top=thin, bottom=thin); cell.alignment = Alignment(vertical="top", wrap_text=True)
        if not has_data: raise ValueError("拒绝生成无数据 Excel")
        workbook.save(path); reopened = load_workbook(path, data_only=False)
        if not reopened.sheetnames or not any(ws.max_row and any(c.value not in (None, "") for row in ws.iter_rows() for c in row) for ws in reopened.worksheets): reopened.close(); path.unlink(missing_ok=True); raise RuntimeError("XLSX 回读为空")
        reopened.close(); return self._result(path, "create_xlsx", sources)

    def create_pdf(self, session_id: str, filename: str, content: str, sources: list[str] | None = None) -> ExportedFile:
        self._require(content)
        try: import fitz
        except ImportError as exc: raise RuntimeError("生成 PDF 需要 PyMuPDF") from exc
        path = self._path(session_id, filename, ".pdf"); document = fitz.open(); page = document.new_page()
        font_path = next((Path(x) for x in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simsun.ttc") if Path(x).is_file()), None)
        try:
            if font_path: page.insert_font(fontname="xiaoche", fontfile=str(font_path)); font = "xiaoche"
            else: font = "china-s"
            rect = fitz.Rect(50, 50, page.rect.width - 50, page.rect.height - 50)
            remaining = str(content)
            while remaining:
                chunk, remaining = remaining[:1500], remaining[1500:]
                result = page.insert_textbox(rect, chunk, fontname=font, fontsize=11, lineheight=1.4)
                if result < 0 and remaining: page = document.new_page()
            document.save(path)
        except Exception: document.close(); path.unlink(missing_ok=True); raise
        document.close(); reopened = fitz.open(path); extracted = "".join(page.get_text() for page in reopened); pages = reopened.page_count; reopened.close()
        if pages < 1 or not extracted.strip(): path.unlink(missing_ok=True); raise RuntimeError("PDF 回读为空")
        return self._result(path, "create_pdf", sources)
